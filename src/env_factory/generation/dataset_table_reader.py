"""Read bounded tabular samples from supported dataset source files."""

from __future__ import annotations

import csv
import gzip
import io
import json
import sqlite3
import tarfile
import tempfile
import zipfile
from contextlib import contextmanager
from itertools import islice
from pathlib import Path
from typing import Any, BinaryIO, Iterator

import ijson

from env_factory.generation.dataset_formats import (
    ARCHIVE_EXTENSIONS, SUPPORTED_EXTENSIONS, SUPPORTED_SOURCE_EXTENSIONS, source_extension,
)
from env_factory.generation.task_generator import TaskGenerationError


DEFAULT_MAX_SOURCE_BYTES = 5_000_000_000
MAX_SAMPLE_ROWS = 500
MAX_ROW_BYTES = 10_000_000


def _safe_member(name: str) -> bool:
    path = Path(name)
    return bool(name) and not path.is_absolute() and ".." not in path.parts and "\\" not in name


def selected_member(path: Path, *, max_source_bytes: int = DEFAULT_MAX_SOURCE_BYTES) -> str | None:
    """Return the first supported table member without extracting an archive."""
    kind = source_extension(path.name)
    if kind not in ARCHIVE_EXTENSIONS:
        return None
    if kind == ".gz":
        member = path.name[:-3]
        if source_extension(member) not in SUPPORTED_EXTENSIONS:
            raise TaskGenerationError("gzip source does not contain a supported table format")
        return member
    if kind == ".zip":
        with zipfile.ZipFile(path) as archive:
            entries = [(item.filename, item.file_size) for item in archive.infolist()
                       if not item.is_dir() and not ((item.external_attr >> 16) & 0o170000 == 0o120000)]
    else:
        with tarfile.open(path, "r:*") as archive:
            entries = [(item.name, item.size) for item in archive.getmembers() if item.isfile()]
    if len(entries) > 10_000:
        raise TaskGenerationError("archive contains too many files")
    if any(not _safe_member(name) for name, _ in entries):
        raise TaskGenerationError("archive contains unsafe member paths")
    candidates = [(name, size) for name, size in entries
                  if source_extension(name) in SUPPORTED_EXTENSIONS and size <= max_source_bytes]
    if not candidates:
        raise TaskGenerationError("archive contains no supported table within the size limit")
    preference = {suffix: rank for rank, suffix in enumerate((
        ".csv", ".tsv", ".jsonl", ".ndjson", ".json", ".xlsx",
        ".parquet", ".sqlite", ".sqlite3", ".db",
    ))}
    return min(candidates, key=lambda item: (preference[source_extension(item[0])], item[0]))[0]


@contextmanager
def _open_member(path: Path, member: str) -> Iterator[BinaryIO]:
    kind = source_extension(path.name)
    if kind == ".gz":
        with gzip.open(path, "rb") as stream:
            yield stream
    elif kind == ".zip":
        with zipfile.ZipFile(path) as archive, archive.open(member) as stream:
            yield stream
    else:
        with tarfile.open(path, "r:*") as archive:
            stream = archive.extractfile(member)
            if stream is None:
                raise TaskGenerationError("archive member cannot be opened")
            with stream:
                yield stream


@contextmanager
def _open_binary(path: Path, member: str | None) -> Iterator[BinaryIO]:
    if member is None:
        with path.open("rb") as stream:
            yield stream
    else:
        with _open_member(path, member) as stream:
            yield stream


@contextmanager
def _open_text(path: Path, member: str | None, encoding: str) -> Iterator[io.TextIOWrapper]:
    with _open_binary(path, member) as binary:
        with io.TextIOWrapper(binary, encoding=encoding, newline="") as stream:
            yield stream


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (dict, list, tuple, bytes)):
        raise TaskGenerationError("dataset table contains nested or binary values")
    if hasattr(value, "isoformat"):
        return str(value.isoformat())
    return str(value)


def _normalize(rows: list[dict[str, Any]], headers: list[Any]) -> tuple[list[str], list[dict[str, str]]]:
    if not all(isinstance(header, str) and header.strip() for header in headers):
        raise TaskGenerationError("dataset table needs non-empty string headers")
    names = [header.strip() for header in headers]
    if len(names) != len(set(names)) or len(names) < 3:
        raise TaskGenerationError("dataset table needs at least three unique headers")
    if len(rows) < 5:
        raise TaskGenerationError("dataset table needs at least five rows")
    if any(set(row) != set(headers) for row in rows):
        raise TaskGenerationError("dataset rows have inconsistent fields")
    return names, [{name: _text(row[original]) for name, original in zip(names, headers)}
                   for row in rows]


def _delimited(path: Path, member: str | None, delimiter: str) -> tuple[list[Any], list[dict[str, Any]]]:
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            with _open_text(path, member, encoding) as stream:
                reader = csv.DictReader(stream, delimiter=delimiter)
                headers = reader.fieldnames or []
                rows = list(islice(reader, MAX_SAMPLE_ROWS))
            return headers, rows
        except UnicodeError:
            continue
    raise TaskGenerationError("dataset text has unsupported encoding")


def _json_rows(path: Path, member: str | None) -> tuple[list[Any], list[dict[str, Any]]]:
    with _open_binary(path, member) as stream:
        first = stream.read(4096).lstrip(b"\xef\xbb\xbf \t\r\n")[:1]
    if first == b"[":
        prefix = "item"
    elif first == b"{":
        with _open_binary(path, member) as stream:
            prefix = next((f"{name}.item" for name, event, _ in ijson.parse(stream)
                           if event == "start_array" and name and "." not in name), None)
        if prefix is None:
            raise TaskGenerationError("JSON source needs an array of row objects")
    else:
        raise TaskGenerationError("JSON source needs an array of row objects")
    with _open_binary(path, member) as stream:
        rows = list(islice(ijson.items(stream, prefix), MAX_SAMPLE_ROWS))
    if not rows or not isinstance(rows[0], dict):
        raise TaskGenerationError("JSON source needs an array of row objects")
    return list(rows[0]), rows


def _json_lines(path: Path, member: str | None) -> tuple[list[Any], list[dict[str, Any]]]:
    with _open_text(path, member, "utf-8-sig") as stream:
        rows = []
        for line in stream:
            if len(line) > MAX_ROW_BYTES:
                raise TaskGenerationError("JSON Lines row exceeds 10 MB")
            if line.strip():
                rows.append(json.loads(line))
            if len(rows) >= MAX_SAMPLE_ROWS:
                break
    if not rows or not isinstance(rows[0], dict):
        raise TaskGenerationError("JSON Lines source needs row objects")
    return list(rows[0]), rows


def _xlsx(path: Path, max_source_bytes: int) -> tuple[list[Any], list[dict[str, Any]]]:
    from openpyxl import load_workbook

    with zipfile.ZipFile(path) as archive:
        if sum(info.file_size for info in archive.infolist()) > max_source_bytes:
            raise TaskGenerationError("Excel source expands beyond the size limit")
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            cells = sheet.iter_rows(values_only=True)
            headers = next(cells, ())
            if not headers or len(headers) < 3:
                continue
            rows = [dict(zip(headers, row)) for row in islice(cells, MAX_SAMPLE_ROWS)]
            if len(rows) >= 5:
                return list(headers), rows
    finally:
        workbook.close()
    raise TaskGenerationError("Excel source has no sheet with three columns and five rows")


def _parquet(path: Path) -> tuple[list[Any], list[dict[str, Any]]]:
    import pyarrow.parquet as pq

    parquet = pq.ParquetFile(path)
    headers = parquet.schema_arrow.names
    rows: list[dict[str, Any]] = []
    for batch in parquet.iter_batches(batch_size=MAX_SAMPLE_ROWS):
        rows.extend(batch.to_pylist()[:MAX_SAMPLE_ROWS - len(rows)])
        if len(rows) >= MAX_SAMPLE_ROWS:
            break
    return headers, rows


def _sqlite(path: Path) -> tuple[list[Any], list[dict[str, Any]]]:
    connection = sqlite3.connect(path.resolve().as_uri() + "?mode=ro", uri=True)
    try:
        names = connection.execute(
            "SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%' ORDER BY name"
        ).fetchall()
        for (name,) in names:
            quoted = '"' + name.replace('"', '""') + '"'
            cursor = connection.execute(f"SELECT * FROM {quoted} LIMIT {MAX_SAMPLE_ROWS}")
            headers = [column[0] for column in cursor.description]
            rows = [dict(zip(headers, row)) for row in cursor]
            if len(headers) >= 3 and len(rows) >= 5:
                return headers, rows
    finally:
        connection.close()
    raise TaskGenerationError("SQLite source has no table with three columns and five rows")


def sample_table(path: Path, *, max_source_bytes: int = DEFAULT_MAX_SOURCE_BYTES) -> tuple[list[str], list[dict[str, str]]]:
    suffix = source_extension(path.name)
    if suffix not in SUPPORTED_SOURCE_EXTENSIONS:
        raise TaskGenerationError(f"unsupported dataset format: {suffix or '<none>'}")
    if path.stat().st_size > max_source_bytes:
        raise TaskGenerationError("dataset source exceeds the configured size limit")
    try:
        member = selected_member(path, max_source_bytes=max_source_bytes)
        table_suffix = source_extension(member) if member else suffix
        if member and table_suffix not in {".csv", ".tsv", ".json", ".jsonl", ".ndjson"}:
            with tempfile.TemporaryDirectory(prefix="envfactory-source-") as temporary:
                extracted = Path(temporary) / ("table" + table_suffix)
                with _open_member(path, member) as source, extracted.open("wb") as destination:
                    copied = 0
                    while block := source.read(1024 * 1024):
                        copied += len(block)
                        if copied > max_source_bytes:
                            raise TaskGenerationError("archive member exceeds the configured size limit")
                        destination.write(block)
                return sample_table(extracted, max_source_bytes=max_source_bytes)
        if table_suffix in {".csv", ".tsv"}:
            headers, rows = _delimited(path, member, "," if table_suffix == ".csv" else "\t")
        elif table_suffix == ".json":
            headers, rows = _json_rows(path, member)
        elif table_suffix in {".jsonl", ".ndjson"}:
            headers, rows = _json_lines(path, member)
        elif table_suffix == ".xlsx":
            headers, rows = _xlsx(path, max_source_bytes)
        elif table_suffix == ".parquet":
            headers, rows = _parquet(path)
        else:
            headers, rows = _sqlite(path)
        return _normalize(rows, headers)
    except (OSError, UnicodeError, ValueError, TypeError, sqlite3.DatabaseError,
            json.JSONDecodeError, zipfile.BadZipFile, tarfile.TarError,
            csv.Error, ijson.JSONError, RuntimeError) as exc:
        raise TaskGenerationError(f"cannot read dataset source {path.name}: {exc}") from exc
